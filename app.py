from __future__ import annotations

import os
import re
import tempfile
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from faster_whisper import WhisperModel
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


APP_NAME = "Moun AI Transcribe"
MODEL_NAME = "base"
MAX_FILE_MB = 200
SUPPORTED_TYPES = ["mp3", "wav", "m4a", "mp4", "webm", "ogg", "flac"]

LANGUAGES = {
    "Tự động nhận diện": None,
    "Tiếng Việt": "vi",
    "Tiếng Lào": "lo",
    "Tiếng Thái": "th",
    "Tiếng Anh": "en",
}


st.set_page_config(
    page_title=APP_NAME,
    page_icon="🎙️",
    layout="wide",
    initial_sidebar_state="expanded",
)


# -----------------------------
# GIAO DIỆN
# -----------------------------
st.markdown(
    """
    <style>
        .stApp {
            background:
                radial-gradient(circle at 0% 0%, rgba(255, 200, 87, 0.18), transparent 28%),
                radial-gradient(circle at 100% 5%, rgba(44, 196, 179, 0.18), transparent 30%),
                linear-gradient(135deg, #fffaf0 0%, #f5fffd 45%, #f6f8ff 100%);
        }
        .block-container {
            max-width: 1180px;
            padding-top: 2rem;
            padding-bottom: 3rem;
        }
        .hero {
            padding: 28px 30px;
            border-radius: 26px;
            background: linear-gradient(135deg, rgba(255,188,71,0.95), rgba(35,184,166,0.95));
            color: white;
            box-shadow: 0 18px 50px rgba(32, 112, 106, 0.18);
            margin-bottom: 22px;
        }
        .hero h1 {
            margin: 0 0 8px 0;
            font-size: clamp(2rem, 5vw, 3.4rem);
            line-height: 1.05;
        }
        .hero p {
            margin: 0;
            font-size: 1.05rem;
            opacity: 0.96;
        }
        .feature-card {
            border: 1px solid rgba(23, 121, 111, 0.12);
            border-radius: 20px;
            padding: 18px 20px;
            background: rgba(255,255,255,0.78);
            box-shadow: 0 10px 30px rgba(34, 80, 75, 0.06);
            min-height: 116px;
        }
        .feature-card b { font-size: 1.02rem; }
        .small-note { color: #5d6c6a; font-size: 0.92rem; }
        div[data-testid="stFileUploader"] {
            background: rgba(255,255,255,0.74);
            border-radius: 20px;
            padding: 10px 14px 4px 14px;
            border: 1px dashed rgba(35,184,166,0.4);
        }
        div[data-testid="stDownloadButton"] button,
        div[data-testid="stButton"] button {
            border-radius: 14px;
            min-height: 44px;
            font-weight: 700;
        }
        textarea {
            border-radius: 16px !important;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="hero">
        <h1>🎙️ Moun AI Transcribe</h1>
        <p>Chuyển âm thanh / video thành văn bản, sau đó tải về TXT, Word, PDF hoặc phụ đề SRT.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

c1, c2, c3 = st.columns(3)
with c1:
    st.markdown(
        '<div class="feature-card"><b>⚡ Nhận diện bằng AI</b><br><span class="small-note">Whisper chạy trên CPU, phù hợp cho bản demo và học tập.</span></div>',
        unsafe_allow_html=True,
    )
with c2:
    st.markdown(
        '<div class="feature-card"><b>🌏 Nhiều ngôn ngữ</b><br><span class="small-note">Tự động nhận diện hoặc chọn Việt, Lào, Thái, Anh.</span></div>',
        unsafe_allow_html=True,
    )
with c3:
    st.markdown(
        '<div class="feature-card"><b>📄 Xuất nhiều định dạng</b><br><span class="small-note">TXT, DOCX, PDF và SRT chỉ bằng một lần chuyển đổi.</span></div>',
        unsafe_allow_html=True,
    )

st.write("")


# -----------------------------
# HÀM XỬ LÝ
# -----------------------------
@st.cache_resource(show_spinner=False)
def load_model() -> WhisperModel:
    """Tải model một lần và dùng lại cho các lần xử lý sau."""
    return WhisperModel(MODEL_NAME, device="cpu", compute_type="int8")


def format_clock(seconds: float) -> str:
    total = max(0, int(seconds))
    hours = total // 3600
    minutes = (total % 3600) // 60
    secs = total % 60
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def format_srt_time(seconds: float) -> str:
    milliseconds = max(0, int(round(seconds * 1000)))
    hours = milliseconds // 3_600_000
    milliseconds %= 3_600_000
    minutes = milliseconds // 60_000
    milliseconds %= 60_000
    secs = milliseconds // 1000
    ms = milliseconds % 1000
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{ms:03d}"


def safe_stem(filename: str) -> str:
    stem = Path(filename).stem.strip() or "transcript"
    stem = re.sub(r"[^\w\- .]+", "_", stem, flags=re.UNICODE)
    return stem[:80].strip(" ._") or "transcript"


def make_docx(text: str, source_name: str, detected_language: str) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("VĂN BẢN CHUYỂN TỪ ÂM THANH / VIDEO")
    run.bold = True
    run.font.size = Pt(18)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"Nguồn: {source_name}  •  Ngôn ngữ: {detected_language}")
    meta_run.italic = True
    meta_run.font.size = Pt(10)

    doc.add_paragraph("")
    for line in text.splitlines():
        if line.strip():
            p = doc.add_paragraph(line.strip())
            p.paragraph_format.space_after = Pt(6)
        else:
            doc.add_paragraph("")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def find_unicode_font() -> str | None:
    candidates = [
        # Windows
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
        r"C:\Windows\Fonts\segoeui.ttf",
        # Linux / Streamlit Cloud
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        # macOS
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for font_path in candidates:
        if os.path.exists(font_path):
            return font_path
    return None


def make_pdf(text: str, source_name: str, detected_language: str) -> bytes:
    font_path = find_unicode_font()
    if not font_path:
        raise RuntimeError(
            "Máy chủ không tìm thấy font Unicode để tạo PDF tiếng Việt/Lào/Thái. "
            "Bạn vẫn có thể tải TXT và Word."
        )

    font_name = "AppUnicodeFont"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, font_path))

    buffer = BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="AI Transcript",
        author=APP_NAME,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "AppTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=17,
        leading=22,
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    meta_style = ParagraphStyle(
        "AppMeta",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        textColor="#5c6664",
        spaceAfter=14,
    )
    body_style = ParagraphStyle(
        "AppBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=11,
        leading=17,
        spaceAfter=7,
    )

    story = [
        Paragraph("VĂN BẢN CHUYỂN TỪ ÂM THANH / VIDEO", title_style),
        Paragraph(
            escape(f"Nguồn: {source_name} • Ngôn ngữ: {detected_language}"),
            meta_style,
        ),
        Spacer(1, 4),
    ]

    for line in text.splitlines():
        if line.strip():
            story.append(Paragraph(escape(line.strip()), body_style))
        else:
            story.append(Spacer(1, 8))

    pdf.build(story)
    return buffer.getvalue()


def make_srt(segments: list[dict]) -> str:
    blocks = []
    for index, segment in enumerate(segments, start=1):
        text = segment["text"].strip()
        if not text:
            continue
        blocks.append(
            f"{index}\n"
            f"{format_srt_time(segment['start'])} --> {format_srt_time(segment['end'])}\n"
            f"{text}"
        )
    return "\n\n".join(blocks) + ("\n" if blocks else "")


def transcribe_file(
    file_path: str,
    language_code: str | None,
    with_timestamps: bool,
    progress_bar,
    status_box,
):
    model = load_model()
    status_box.write("🧠 Đang nhận diện giọng nói...")

    segments_generator, info = model.transcribe(
        file_path,
        language=language_code,
        beam_size=5,
        vad_filter=True,
        condition_on_previous_text=True,
    )

    collected = []
    text_lines = []
    duration = max(float(info.duration or 0), 1.0)

    for segment in segments_generator:
        segment_text = segment.text.strip()
        if not segment_text:
            continue

        item = {
            "start": float(segment.start),
            "end": float(segment.end),
            "text": segment_text,
        }
        collected.append(item)

        if with_timestamps:
            text_lines.append(f"[{format_clock(segment.start)}] {segment_text}")
        else:
            text_lines.append(segment_text)

        percent = min(int((float(segment.end) / duration) * 100), 99)
        progress_bar.progress(percent)

    progress_bar.progress(100)
    status_box.write("✅ Hoàn thành nhận diện.")

    detected = str(info.language or "unknown")
    probability = float(info.language_probability or 0.0)
    return "\n".join(text_lines), collected, detected, probability


# -----------------------------
# SIDEBAR
# -----------------------------
with st.sidebar:
    st.header("⚙️ Cài đặt")
    language_label = st.selectbox(
        "Ngôn ngữ trong file",
        options=list(LANGUAGES.keys()),
        index=0,
        help="Nếu không chắc, để Tự động nhận diện.",
    )
    with_timestamps = st.toggle(
        "Thêm thời gian vào văn bản",
        value=False,
        help="Ví dụ: [00:01:12] Nội dung...",
    )

    st.divider()
    st.caption(f"Model: Whisper `{MODEL_NAME}` • CPU int8")
    st.caption(f"Giới hạn upload cấu hình: {MAX_FILE_MB} MB")
    st.info(
        "File được ghi tạm trong lúc xử lý và code sẽ xóa file tạm sau khi hoàn tất. "
        "Ứng dụng này không tự lưu lịch sử transcript vào database."
    )


# -----------------------------
# UPLOAD + XỬ LÝ
# -----------------------------
st.subheader("1. Chọn file")
uploaded_file = st.file_uploader(
    "Kéo thả file vào đây hoặc bấm Browse files",
    type=SUPPORTED_TYPES,
    help="Hỗ trợ MP3, WAV, M4A, MP4, WEBM, OGG, FLAC.",
)

if uploaded_file is None:
    st.markdown(
        "**Cách dùng:** tải file lên → chọn ngôn ngữ → bấm **Chuyển thành văn bản** → chỉnh lại nội dung nếu cần → tải TXT / Word / PDF / SRT."
    )
    st.stop()

file_size_mb = uploaded_file.size / (1024 * 1024)
if file_size_mb > MAX_FILE_MB:
    st.error(f"File lớn hơn {MAX_FILE_MB} MB. Hãy dùng file nhỏ hơn.")
    st.stop()

signature = f"{uploaded_file.name}:{uploaded_file.size}"
if st.session_state.get("upload_signature") != signature:
    st.session_state["upload_signature"] = signature
    st.session_state.pop("transcript_text", None)
    st.session_state.pop("edited_text", None)
    st.session_state.pop("segments", None)
    st.session_state.pop("detected_language", None)
    st.session_state.pop("language_probability", None)

info_col1, info_col2 = st.columns([2, 1])
with info_col1:
    st.write(f"**File:** `{uploaded_file.name}`")
with info_col2:
    st.write(f"**Dung lượng:** {file_size_mb:.2f} MB")

suffix = Path(uploaded_file.name).suffix.lower()
if suffix in {".mp4", ".webm"}:
    st.video(uploaded_file)
else:
    st.audio(uploaded_file)

st.subheader("2. Chuyển thành văn bản")
button_col, clear_col = st.columns([3, 1])
with button_col:
    start = st.button(
        "✨ Chuyển thành văn bản",
        type="primary",
        use_container_width=True,
    )
with clear_col:
    if st.button("🗑️ Xóa kết quả", use_container_width=True):
        for key in [
            "transcript_text",
            "edited_text",
            "segments",
            "detected_language",
            "language_probability",
        ]:
            st.session_state.pop(key, None)
        st.rerun()

if start:
    temp_path = None
    try:
        progress = st.progress(0)
        status = st.empty()
        status.write("📦 Đang chuẩn bị file...")

        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix or ".media") as tmp:
            tmp.write(uploaded_file.getbuffer())
            temp_path = tmp.name

        status.write("🤖 Đang tải model AI (lần đầu có thể lâu hơn)...")
        transcript, segments, detected, probability = transcribe_file(
            temp_path,
            LANGUAGES[language_label],
            with_timestamps,
            progress,
            status,
        )

        if not transcript.strip():
            st.warning("Không nhận diện được lời nói rõ ràng trong file này.")
        else:
            st.session_state["transcript_text"] = transcript
            st.session_state["edited_text"] = transcript
            st.session_state["segments"] = segments
            st.session_state["detected_language"] = detected
            st.session_state["language_probability"] = probability
            st.success("Chuyển đổi thành công!")

    except Exception as exc:
        st.error("Không thể xử lý file. Chi tiết lỗi:")
        st.code(str(exc))
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


# -----------------------------
# KẾT QUẢ + XUẤT FILE
# -----------------------------
if st.session_state.get("transcript_text"):
    detected_language = st.session_state.get("detected_language", "unknown")
    probability = st.session_state.get("language_probability", 0.0)
    segments = st.session_state.get("segments", [])

    st.subheader("3. Kiểm tra và chỉnh sửa")
    st.caption(
        f"AI nhận diện ngôn ngữ: **{detected_language}** "
        f"(độ tin cậy khoảng {probability * 100:.1f}%). Bạn có thể sửa trực tiếp nội dung bên dưới trước khi tải."
    )

    edited_text = st.text_area(
        "Nội dung",
        key="edited_text",
        height=430,
        label_visibility="collapsed",
    )

    st.subheader("4. Tải kết quả")
    base_name = safe_stem(uploaded_file.name)

    txt_bytes = edited_text.encode("utf-8")
    docx_bytes = make_docx(edited_text, uploaded_file.name, detected_language)
    srt_text = make_srt(segments)

    pdf_bytes = None
    pdf_error = None
    try:
        pdf_bytes = make_pdf(edited_text, uploaded_file.name, detected_language)
    except Exception as exc:
        pdf_error = str(exc)

    d1, d2, d3, d4 = st.columns(4)
    with d1:
        st.download_button(
            "⬇️ TXT",
            data=txt_bytes,
            file_name=f"{base_name}.txt",
            mime="text/plain; charset=utf-8",
            use_container_width=True,
        )
    with d2:
        st.download_button(
            "⬇️ Word",
            data=docx_bytes,
            file_name=f"{base_name}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )
    with d3:
        if pdf_bytes is not None:
            st.download_button(
                "⬇️ PDF",
                data=pdf_bytes,
                file_name=f"{base_name}.pdf",
                mime="application/pdf",
                use_container_width=True,
            )
        else:
            st.button("PDF chưa khả dụng", disabled=True, use_container_width=True)
    with d4:
        st.download_button(
            "⬇️ SRT",
            data=srt_text.encode("utf-8"),
            file_name=f"{base_name}.srt",
            mime="application/x-subrip",
            use_container_width=True,
            disabled=not bool(srt_text.strip()),
        )

    if pdf_error:
        st.warning(pdf_error)

    with st.expander("Xem các đoạn có timestamp"):
        if segments:
            for item in segments:
                st.write(
                    f"`{format_clock(item['start'])} → {format_clock(item['end'])}`  {item['text']}"
                )
        else:
            st.caption("Chưa có dữ liệu đoạn.")

st.divider()
st.caption(
    "Moun AI Transcribe • faster-whisper • Khi triển khai trên hosting miễn phí, file dài có thể xử lý chậm do giới hạn CPU/RAM."
)
